from __future__ import annotations

import logging
import threading
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

logger = logging.getLogger(__name__)

# ---- 硬性上限：防止畸形文件拖垮进程 ----
max_file_size_bytes = 30 * 1024 * 1024  # 单份简历最大 30MB
min_file_size_bytes = 1  # 空文件直接拒绝
max_pages = 100  # PDF 最多解析页数
max_chars = 60_000  # 提取文本字符上限，超出截断
parse_timeout_seconds = 30.0  # 单份简历解析超时
pdf_magic = b"%PDF-"
pdf_magic_scan_bytes = 1024  # PDF 规范允许头部前 1024 字节内出现魔数


class ResumeError(Exception):
    """简历解析相关异常的基类。"""


class ResumeFileError(ResumeError):
    """文件本身有问题：不存在、不是文件、为空、过大、不可读。"""


class ResumeUnsupportedError(ResumeError):
    """格式不支持，或后缀与实际内容不符。"""


class ResumeDependencyError(ResumeError):
    """缺少解析依赖，附带安装提示。"""


class ResumeParseError(ResumeError):
    """解析过程失败。"""


class ResumeEmptyContentError(ResumeParseError):
    """解析成功但拿不到任何文本，通常是扫描件或图片型简历。"""


class ResumeTimeoutError(ResumeParseError):
    """解析超时。"""


@dataclass
class ResumeParseResult:
    """统一的解析结果，附带足够信息供上层判断简历质量与是否降级。"""

    text: str
    file_type: str
    source: str
    page_count: int = 0
    char_count: int = 0
    truncated: bool = False
    warnings: list[str] = field(default_factory=list)

    def is_empty(self) -> bool:
        return not self.text.strip()

    def to_dict(self) -> dict[str, Any]:
        return {
            "text": self.text,
            "file_type": self.file_type,
            "source": self.source,
            "page_count": self.page_count,
            "char_count": self.char_count,
            "truncated": self.truncated,
            "warnings": list(self.warnings),
        }


class ResumeAnalyzer:
    """简历解析器，支持 PDF 与 DOCX。"""

    def __init__(
        self,
        resume_file: str | Path,
        *,
        max_file_size: int = max_file_size_bytes,
        max_chars: int = max_chars,
        max_pages: int = max_pages,
        timeout: float = parse_timeout_seconds,
    ) -> None:
        try:
            self.resume_file = Path(resume_file)
        except TypeError as exc:
            raise ResumeFileError(f"非法的简历路径: {resume_file!r}") from exc

        self.max_file_size = max_file_size
        self.max_chars = max_chars
        self.max_pages = max_pages
        self.timeout = timeout

        self._validate_file()

    # ---------- 文件校验 ----------

    def _validate_file(self) -> None:
        """在解析前把文件层面的问题一次性拦掉。"""
        path = self.resume_file

        if not path.exists():
            raise ResumeFileError(f"简历文件不存在: {path}")
        if not path.is_file():
            raise ResumeFileError(f"简历路径不是文件: {path}")

        try:
            size = path.stat().st_size
        except OSError as exc:
            raise ResumeFileError(f"无法读取简历文件信息: {path}") from exc

        if size < min_file_size_bytes:
            raise ResumeFileError(f"简历文件为空: {path}")
        if size > self.max_file_size:
            raise ResumeFileError(
                f"简历文件过大: {size / 1024 / 1024:.1f}MB，"
                f"上限 {self.max_file_size / 1024 / 1024:.0f}MB"
            )

        try:
            with path.open("rb") as fp:
                fp.read(1)
        except OSError as exc:
            raise ResumeFileError(f"简历文件不可读: {path}") from exc

    def select(self) -> str:
        """按后缀判断简历格式，仅做初判，真实格式在解析时校验。"""
        suffix = self.resume_file.suffix.lower()
        if suffix == ".pdf":
            return "pdf"
        if suffix == ".docx":
            return "docx"
        return "unknown"

    # ---------- 对外入口 ----------

    def analyze(self) -> ResumeParseResult:
        """解析简历并返回结构化结果。"""
        file_type = self.select()
        if file_type == "unknown":
            suffix = self.resume_file.suffix or "(无后缀)"
            raise ResumeUnsupportedError(f"不支持的简历格式: {suffix}")

        started = time.monotonic()
        parser = self._analyze_pdf if file_type == "pdf" else self._analyze_docx
        result = self._run_with_timeout(parser)

        logger.info(
            "简历解析完成: type=%s pages=%s chars=%s truncated=%s warnings=%s cost=%.2fs",
            result.file_type,
            result.page_count,
            result.char_count,
            result.truncated,
            len(result.warnings),
            time.monotonic() - started,
        )
        return result

    def analyze_text(self) -> str:
        """仅取纯文本的便捷入口。"""
        return self.analyze().text

    # ---------- 超时与截断 ----------

    def _run_with_timeout(self, func: Callable[[], ResumeParseResult]) -> ResumeParseResult:
        """在守护线程中执行解析并施加超时。

        局限：超时后线程无法被强制终止，仍会在后台跑完。
        因此使用 daemon 线程，保证它不会阻塞进程退出。
        """
        box: dict[str, Any] = {}

        def runner() -> None:
            try:
                box["result"] = func()
            except BaseException as exc:  # noqa: BLE001 - 需完整回传，避免线程内异常被吞
                box["error"] = exc

        worker = threading.Thread(target=runner, name="resume-parse", daemon=True)
        worker.start()
        worker.join(self.timeout)

        if worker.is_alive():
            raise ResumeTimeoutError(
                f"简历解析超时（>{self.timeout:.0f}s）: {self.resume_file}"
            )
        if "error" in box:
            raise box["error"]
        return box["result"]

    def _finalize(
        self, text: str, file_type: str, warnings: list[str], page_count: int = 0
    ) -> ResumeParseResult:
        """收敛结果：空内容兜底、超长截断、统一封装。"""
        text = text.strip()

        if not text:
            raise ResumeEmptyContentError(
                f"未从简历中提取到任何文本: {self.resume_file}。"
                "若为扫描件或图片型简历，请先做 OCR 或改用文本版简历。"
            )

        truncated = False
        if len(text) > self.max_chars:
            warnings.append(
                f"简历文本超过上限 {self.max_chars} 字符，已截断"
                f"（原始 {len(text)} 字符）"
            )
            text = text[: self.max_chars]
            truncated = True

        for item in warnings:
            logger.warning("简历解析提示: %s", item)

        return ResumeParseResult(
            text=text,
            file_type=file_type,
            source=str(self.resume_file),
            page_count=page_count,
            char_count=len(text),
            truncated=truncated,
            warnings=warnings,
        )

    # ---------- PDF ----------

    @staticmethod
    def _import_fitz() -> Any:
        try:
            import fitz  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ResumeDependencyError(
                "解析 PDF 简历需要 PyMuPDF，请执行: pip install pymupdf"
            ) from exc
        return fitz

    def _check_pdf_magic(self) -> None:
        try:
            with self.resume_file.open("rb") as fp:
                head = fp.read(pdf_magic_scan_bytes)
        except OSError as exc:
            raise ResumeFileError(f"简历文件不可读: {self.resume_file}") from exc

        if pdf_magic not in head:
            raise ResumeUnsupportedError(
                f"文件后缀为 .pdf，但内容不是有效 PDF: {self.resume_file}"
            )

    def _analyze_pdf(self) -> ResumeParseResult:
        fitz = self._import_fitz()
        self._check_pdf_magic()
        warnings: list[str] = []

        try:
            doc = fitz.open(str(self.resume_file))
        except Exception as exc:
            raise ResumeParseError(f"PDF 打开失败: {self.resume_file}") from exc

        try:
            # 加密 PDF：先尝试空密码，失败则明确报错而非返回空文本
            if getattr(doc, "needs_pass", False):
                if not doc.authenticate(""):
                    raise ResumeParseError(
                        f"PDF 已加密，无法解析，请提供未加密的简历: {self.resume_file}"
                    )
                warnings.append("PDF 使用空密码解密后解析")

            if getattr(doc, "is_repaired", False):
                warnings.append("PDF 结构损坏，已自动修复后解析，结果可能不完整")

            page_count = int(getattr(doc, "page_count", 0) or 0)
            limit = min(page_count, self.max_pages)
            if page_count > self.max_pages:
                warnings.append(
                    f"PDF 共 {page_count} 页，超过上限 {self.max_pages} 页，"
                    f"仅解析前 {self.max_pages} 页"
                )

            pages: list[str] = []
            for index in range(limit):
                # 单页失败不影响整体，记录警告后继续
                try:
                    text = self._extract_pdf_page(doc[index])
                except Exception as exc:
                    warnings.append(f"第 {index + 1} 页解析失败: {exc}")
                    continue

                if text:
                    pages.append(text)
                else:
                    warnings.append(f"第 {index + 1} 页无可提取文本（可能是图片页）")
        finally:
            try:
                doc.close()
            except Exception:  # noqa: BLE001 - 关闭失败不应覆盖真正的业务异常
                logger.warning("关闭 PDF 文档失败: %s", self.resume_file)

        return self._finalize("\n\n".join(pages), "pdf", warnings, page_count=page_count)

    @staticmethod
    def _extract_pdf_page(page: Any) -> str:
        """单页多策略提取：优先按阅读顺序取文本，失败则退化为按块拼接。"""
        text = ""
        try:
            text = (page.get_text("text", sort=True) or "").strip()
        except Exception:  # noqa: BLE001 - 退化为下一种策略
            text = ""

        if text:
            return text

        # 兜底策略：按文本块提取，块之间用换行拼接
        try:
            blocks = page.get_text("blocks") or []
        except Exception:  # noqa: BLE001 - 无更多策略，返回空
            return ""

        parts = [
            block[4].strip()
            for block in blocks
            if len(block) > 4 and isinstance(block[4], str) and block[4].strip()
        ]
        return "\n".join(parts).strip()

    # ---------- DOCX ----------

    @staticmethod
    def _import_docx() -> Any:
        try:
            import docx  # type: ignore[import-not-found]
        except ImportError as exc:
            raise ResumeDependencyError(
                "解析 DOCX 简历需要 python-docx，请执行: pip install python-docx"
            ) from exc
        return docx

    def _check_docx_magic(self) -> None:
        """DOCX 本质是 zip，用压缩包结构做真实性校验，拦截伪装后缀。"""
        try:
            with zipfile.ZipFile(self.resume_file) as zf:
                names = zf.namelist()
        except zipfile.BadZipFile as exc:
            raise ResumeUnsupportedError(
                f"文件后缀为 .docx，但内容不是有效的 DOCX: {self.resume_file}"
            ) from exc

        if "word/document.xml" not in names:
            raise ResumeUnsupportedError(
                f"DOCX 缺少 word/document.xml，可能是 .doc 等旧格式: {self.resume_file}"
            )

    def _analyze_docx(self) -> ResumeParseResult:
        docx = self._import_docx()
        self._check_docx_magic()
        warnings: list[str] = []

        try:
            document = docx.Document(str(self.resume_file))
        except Exception as exc:
            raise ResumeParseError(f"DOCX 打开失败: {self.resume_file}") from exc

        parts: list[str] = []

        # 正文段落
        try:
            for paragraph in document.paragraphs:
                text = (paragraph.text or "").strip()
                if text:
                    parts.append(text)
        except Exception as exc:
            warnings.append(f"段落解析异常: {exc}")

        # 表格内容（简历里的经历/技能常放在表格中）
        try:
            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text and cell.text.strip()
                    ]
                    if cells:
                        parts.append(" | ".join(cells))
        except Exception as exc:
            warnings.append(f"表格解析异常: {exc}")

        return self._finalize("\n".join(parts), "docx", warnings)
